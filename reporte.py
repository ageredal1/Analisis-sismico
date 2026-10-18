"""
reporte.py
Genera el reporte de resultados en formato .xlsx.
"""

import io
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

import graficos as graf


# ─────────────────────────────────────────────────────────────────────────────
# ESTILOS
# ─────────────────────────────────────────────────────────────────────────────

AZUL_OSCURO  = "FF1A237E"  # encabezado principal
AZUL_MEDIO   = "FF283593"  # sub-encabezado
AZUL_CLARO   = "FFE8EAF6"  # filas pares
VERDE_TITULO = "FF1B5E20"
BORDE_FINO   = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _header(ws, fila, col_ini, col_fin, texto, color_hex=AZUL_OSCURO):
    ws.merge_cells(
        start_row=fila, start_column=col_ini,
        end_row=fila,   end_column=col_fin
    )
    c = ws.cell(row=fila, column=col_ini, value=texto)
    c.font = Font(bold=True, color="FFFFFFFF", size=11)
    c.fill = PatternFill("solid", fgColor=color_hex)
    c.alignment = Alignment(horizontal="center", vertical="center")
    return c


def _celda(ws, fila, col, valor, negrita=False, centrado=True, color_hex=None):
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = Font(bold=negrita, size=10)
    c.alignment = Alignment(horizontal="center" if centrado else "left",
                             vertical="center")
    c.border = BORDE_FINO
    if color_hex:
        c.fill = PatternFill("solid", fgColor=color_hex)
    return c


def _auto_ancho(ws, extra=4):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + extra, 30)


def _crear_imagen_excel(fig, dpi=110, ancho_px=None, alto_px=None):
    """Convierte una figura de matplotlib en un objeto Image de openpyxl."""
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    img = Image(buf)
    if ancho_px and alto_px:
        img.width = ancho_px
        img.height = alto_px
    return img


# ─────────────────────────────────────────────────────────────────────────────
# FUNCIÓN PRINCIPAL
# ─────────────────────────────────────────────────────────────────────────────

def generar_reporte(res: dict, hi: float, Z: float, U: float,
                    S: float, R: float, T_modos) -> bytes:
    """
    Genera el libro Excel con todos los resultados numéricos y gráficos,
    devolviéndolo en formato bytes.
    """
    wb = Workbook()
    wb.remove(wb.active)   # quitar hoja vacía por defecto

    n_pisos = res["n_pisos"]
    n_modos = res["n_modos"]
    pisos_labels = [f"Piso {i+1}" for i in range(n_pisos)]

    # ── Hoja 1: Parámetros de entrada ────────────────────────────────────────
    ws1 = wb.create_sheet("Parámetros de Entrada")
    ws1.sheet_view.showGridLines = True

    _header(ws1, 1, 1, 3, "PARÁMETROS DE ENTRADA")
    params = [
        ("Altura de entrepiso hi",      f"{hi} m"),
        ("Factor de zona sísmica Z",    Z),
        ("Factor de uso U",             U),
        ("Factor de amplif. suelo S",   S),
        ("Factor de reducción R",       R),
        ("Periodo Modo 1 T1",           f"{T_modos[0]} s"),
        ("Periodo Modo 2 T2",           f"{T_modos[1]} s"),
        ("Periodo Modo 3 T3",           f"{T_modos[2]} s"),
    ]
    for i, (nombre, valor) in enumerate(params, start=2):
        color = AZUL_CLARO if i % 2 == 0 else None
        _celda(ws1, i, 1, nombre, centrado=False, color_hex=color)
        _celda(ws1, i, 2, valor,  centrado=True,  color_hex=color)
        _celda(ws1, i, 3, "",     centrado=True,  color_hex=color)
    _auto_ancho(ws1)

    # ── Hoja 2: Parámetros Modales ───────────────────────────────────────────
    ws2 = wb.create_sheet("Parámetros Modales")
    ws2.sheet_view.showGridLines = True

    _header(ws2, 1, 1, 5, "PERÍODO FUNDAMENTAL DE VIBRACIÓN APROXIMADO")
    _celda(ws2, 2, 1, "T_aprox (s)", negrita=True, color_hex=AZUL_CLARO)
    _celda(ws2, 2, 2, round(float(res["T_aprox"]), 6))

    fila = 4
    _header(ws2, fila, 1, 6, "PARÁMETROS MODALES POR MODO")
    fila += 1
    encabezados = ["Modo", "Período T (s)", "C", "Sa (m/s²)", "W (rad/s)", "Sd (m)"]
    for j, enc in enumerate(encabezados, start=1):
        _celda(ws2, fila, j, enc, negrita=True, color_hex="FFB3C6E7")
    for modo in range(n_modos):
        fila += 1
        color = AZUL_CLARO if modo % 2 == 0 else None
        vals = [
            f"Modo {modo+1}",
            round(float(T_modos[modo]), 6),
            round(float(res["C"][modo]),  6),
            round(float(res["Sa"][modo]), 6),
            round(float(res["W"][modo]),  6),
            round(float(res["Sd"][modo]), 6),
        ]
        for j, v in enumerate(vals, start=1):
            _celda(ws2, fila, j, v, color_hex=color)
    _auto_ancho(ws2)

    # ── Hoja 3: Fuerzas y Cortantes por Modo ─────────────────────────────────
    ws3 = wb.create_sheet("Fuerzas y Cortantes")
    ws3.sheet_view.showGridLines = True

    _header(ws3, 1, 1, 7, "FUERZA SÍSMICA Y CORTANTE POR PISO Y MODO")
    fila = 2
    cols_enc = ["Piso", "F Modo1 (t)", "F Modo2 (t)", "F Modo3 (t)",
                "V Modo1 (t)", "V Modo2 (t)", "V Modo3 (t)"]
    for j, enc in enumerate(cols_enc, start=1):
        _celda(ws3, fila, j, enc, negrita=True, color_hex="FFB3C6E7")
    for i in range(n_pisos):
        fila += 1
        color = AZUL_CLARO if i % 2 == 0 else None
        _celda(ws3, fila, 1, pisos_labels[i], color_hex=color)
        for modo in range(n_modos):
            _celda(ws3, fila, 2 + modo, round(float(res["F"][i, modo]), 6), color_hex=color)
            _celda(ws3, fila, 5 + modo, round(float(res["V"][i, modo]), 6), color_hex=color)
    _auto_ancho(ws3)

    # Insertar Gráfico 1 en Hoja 3 (al lado de la tabla, columna I)
    try:
        fig_fc = graf.grafico_fuerza_cortante(res["F"], res["V"], n_pisos, n_modos)
        img_fc = _crear_imagen_excel(fig_fc, dpi=110, ancho_px=750, alto_px=340)
        ws3.add_image(img_fc, "I2")
    except Exception as e:
        print(f"Error al adjuntar gráfico a Hoja 3: {e}")

    # ── Hoja 4: Cortantes de Combinación ─────────────────────────────────────
    ws4 = wb.create_sheet("Cortantes Combinación")
    ws4.sheet_view.showGridLines = True

    _header(ws4, 1, 1, 5, "CORTANTES DE COMBINACIÓN MODAL POR PISO")
    fila = 2
    cols_enc = ["Piso", "Vsum_abs (t)", "Vrcsc (t)", "Vrnc_h (t)", "Vreal (t)"]
    for j, enc in enumerate(cols_enc, start=1):
        _celda(ws4, fila, j, enc, negrita=True, color_hex="FFB3C6E7")
    for i in range(n_pisos):
        fila += 1
        color = AZUL_CLARO if i % 2 == 0 else None
        _celda(ws4, fila, 1, pisos_labels[i], color_hex=color)
        _celda(ws4, fila, 2, round(float(res["Vsum_abs"][i]), 6), color_hex=color)
        _celda(ws4, fila, 3, round(float(res["Vrcsc"][i]),    6), color_hex=color)
        _celda(ws4, fila, 4, round(float(res["Vrnc_h"][i]),   6), color_hex=color)
        _celda(ws4, fila, 5, round(float(res["Vreal"][i]),    6), color_hex=color)
    _auto_ancho(ws4)

    # Insertar Gráfico 2 en Hoja 4 (al lado de la tabla, columna G)
    try:
        fig_comb = graf.grafico_cortantes_combinacion(
            res["Vsum_abs"], res["Vrcsc"], res["Vrnc_h"], res["Vreal"], n_pisos
        )
        img_comb = _crear_imagen_excel(fig_comb, dpi=110, ancho_px=750, alto_px=560)
        ws4.add_image(img_comb, "G2")
    except Exception as e:
        print(f"Error al adjuntar gráfico a Hoja 4: {e}")

    # ── Hoja 5: Resumen de Gráficos Dedicada ──────────────────────────────────
    ws5 = wb.create_sheet("Gráficos")
    ws5.sheet_view.showGridLines = True
    _header(ws5, 1, 2, 14, "ANÁLISIS SÍSMICO MODAL – RESUMEN GRÁFICO")

    try:
        fig_g1 = graf.grafico_fuerza_cortante(res["F"], res["V"], n_pisos, n_modos)
        img_g1 = _crear_imagen_excel(fig_g1, dpi=110, ancho_px=800, alto_px=360)
        ws5.add_image(img_g1, "B3")

        fig_g2 = graf.grafico_cortantes_combinacion(
            res["Vsum_abs"], res["Vrcsc"], res["Vrnc_h"], res["Vreal"], n_pisos
        )
        img_g2 = _crear_imagen_excel(fig_g2, dpi=110, ancho_px=800, alto_px=600)
        ws5.add_image(img_g2, "B23")
    except Exception as e:
        print(f"Error al crear hoja de gráficos: {e}")

    # ── Guardar en bytes ──────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
