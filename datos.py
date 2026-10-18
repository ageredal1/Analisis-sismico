"""
datos.py
Módulo para el manejo y lectura de datos de entrada en Excel,
así como la generación en memoria de la plantilla .xlsx.
"""

import io
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ─────────────────────────────────────────────────────────────────────────────
# GENERACIÓN DE PLANTILLA EXCEL
# ─────────────────────────────────────────────────────────────────────────────

def generar_plantilla_bytes() -> bytes:
    """
    Genera la plantilla Excel en memoria con formato profesional y datos
    de ejemplo tomados de ET2_logic.py (5 pisos, masas y matriz de modos Phi).

    Retorna
    -------
    bytes
        Contenido binario del archivo .xlsx listo para descargar.
    """
    AZUL_OSC = "1A237E"
    AZUL_MED = "283593"
    AZUL_CLA = "E8EAF6"
    borde = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    # Cabeceras en la Fila 1 para compatibilidad directa con lectura tabular
    headers = ["Pisos", "Masas", "Modo 1", "Modo 2", "Modo 3"]
    anchos  = [12, 14, 14, 14, 14]
    for j, (h, w) in enumerate(zip(headers, anchos), start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=AZUL_OSC)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borde
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 24

    # Datos de ejemplo (ET2_logic.py)
    masas = [31.80, 31.80, 31.80, 31.80, 27.50]
    phi = [
        [0.03112, -0.08102,  0.10415],
        [0.05959, -0.10193,  0.03021],
        [0.08302, -0.05484, -0.09525],
        [0.09940,  0.03358, -0.05769],
        [0.10834,  0.10609,  0.09176],
    ]
    for i in range(5):
        fila = i + 2
        fill = PatternFill("solid", fgColor=AZUL_CLA) if i % 2 == 0 else None
        for j, val in enumerate([f"Piso {i+1}", masas[i]] + phi[i], start=1):
            c = ws.cell(row=fila, column=j, value=val)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = borde
            if fill:
                c.fill = fill
        ws.row_dimensions[fila].height = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# LECTURA DE ARCHIVO EXCEL DE USUARIO
# ─────────────────────────────────────────────────────────────────────────────

def leer_excel_usuario(archivo) -> tuple[np.ndarray, np.ndarray]:
    """
    Lee las masas de cada piso y la matriz modal Phi desde el archivo .xlsx
    subido por el usuario, detectando automáticamente la fila de encabezados.

    Parámetros
    ----------
    archivo : UploadedFile o ruta o BytesIO
        Archivo Excel a leer.

    Retorna
    -------
    tuple[np.ndarray, np.ndarray]
        (m, Phi) donde m es un vector (n_pisos,) y Phi es una matriz (n_pisos, 3).
    """
    xls = pd.ExcelFile(archivo)
    nombre_hoja = "Datos" if "Datos" in xls.sheet_names else xls.sheet_names[0]
    df = pd.read_excel(xls, sheet_name=nombre_hoja, header=None)

    # Buscar en qué fila se encuentran los encabezados (buscar la palabra 'masa')
    header_row_idx = None
    for r_idx in range(min(10, len(df))):
        row_values = [str(val).strip().lower() for val in df.iloc[r_idx].dropna()]
        if any("masa" in val for val in row_values):
            header_row_idx = r_idx
            break

    if header_row_idx is not None:
        df.columns = [str(c).strip() for c in df.iloc[header_row_idx]]
        df = df.iloc[header_row_idx + 1:].reset_index(drop=True)
    else:
        df.columns = [str(c).strip() for c in df.iloc[0]]
        df = df.iloc[1:].reset_index(drop=True)

    # Identificar columna de masas
    col_masa = next(
        (c for c in df.columns if "masa" in c.lower()),
        None
    )
    if col_masa is None:
        raise KeyError("No se encontró la columna 'Masas' en el archivo.")

    # Identificar columnas de Modos (Modo 1, Modo 2, Modo 3)
    cols_modo = []
    for num in [1, 2, 3]:
        c_found = next(
            (c for c in df.columns if "modo" in c.lower() and str(num) in c),
            None
        )
        if c_found:
            cols_modo.append(c_found)

    # Si no encontró por nombre explícito con número, buscar columnas que contengan 'modo'
    if len(cols_modo) < 3:
        cols_modo = [c for c in df.columns if "modo" in c.lower()]

    # Si aún no, tomar las 3 columnas después de masas
    if len(cols_modo) < 3:
        idx_m = list(df.columns).index(col_masa)
        cols_modo = list(df.columns)[idx_m + 1 : idx_m + 4]

    # Limpiar filas vacías
    df = df.dropna(subset=[col_masa] + cols_modo)

    m = pd.to_numeric(df[col_masa], errors="coerce").to_numpy(dtype=float)
    Phi = df[cols_modo].apply(pd.to_numeric, errors="coerce").to_numpy(dtype=float)

    return m, Phi
